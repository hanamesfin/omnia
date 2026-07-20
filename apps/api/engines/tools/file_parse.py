"""
Extract text from uploaded files — PDF, DOCX, CSV, XLSX, plain text.
"""
from __future__ import annotations

import csv
import io
import json
from typing import Any


def parse_bytes(raw: bytes, filename: str, media: str, *, cap: int = 60_000) -> str:
    ext = (filename.rsplit(".", 1)[-1].lower() if "." in filename else "")

    if ext in ("txt", "md", "markdown", "json", "jsonl", "log", "sql", "py", "js", "ts", "html", "xml", "yaml", "yml"):
        return _decode_text(raw, cap)

    if ext in ("csv", "tsv"):
        return _parse_csv(raw, delimiter="\t" if ext == "tsv" else ",", cap=cap)

    if ext == "pdf" or media == "pdf":
        return _parse_pdf(raw, filename, cap)

    if ext == "docx":
        return _parse_docx(raw, filename, cap)

    if ext in ("xlsx", "xls"):
        return _parse_xlsx(raw, filename, cap)

    if media == "text" or media == "table":
        return _decode_text(raw, cap)

    return f"[Unsupported parse type for {filename} ({len(raw):,} bytes). Supported: pdf, docx, csv, xlsx, text.]"


def _truncate(text: str, cap: int) -> str:
    if len(text) <= cap:
        return text
    return text[:cap] + f"\n\n… [truncated at {cap:,} chars]"


def _decode_text(raw: bytes, cap: int) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            return _truncate(raw.decode(enc), cap)
        except UnicodeDecodeError:
            continue
    return f"[Could not decode as text — {len(raw):,} bytes]"


def _parse_csv(raw: bytes, *, delimiter: str, cap: int) -> str:
    text = _decode_text(raw, cap * 2)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)[:500]
    return _truncate(json.dumps(rows, ensure_ascii=False, indent=0), cap)


def _parse_pdf(raw: bytes, filename: str, cap: int) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        return (
            f"[PDF {filename}: install pypdf on the API host to extract text. "
            f"{len(raw):,} bytes stored.]"
        )
    try:
        reader = PdfReader(io.BytesIO(raw))
        parts: list[str] = []
        for page in reader.pages[:40]:
            parts.append(page.extract_text() or "")
        return _truncate("\n\n".join(parts).strip(), cap)
    except Exception as e:
        return f"[PDF parse failed for {filename}: {e}]"


def _parse_docx(raw: bytes, filename: str, cap: int) -> str:
    try:
        from docx import Document  # type: ignore
    except ImportError:
        return (
            f"[DOCX {filename}: install python-docx on the API host. "
            f"{len(raw):,} bytes stored.]"
        )
    try:
        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        return _truncate("\n".join(parts), cap)
    except Exception as e:
        return f"[DOCX parse failed for {filename}: {e}]"


def _parse_xlsx(raw: bytes, filename: str, cap: int) -> str:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        return (
            f"[XLSX {filename}: install openpyxl on the API host. "
            f"{len(raw):,} bytes stored.]"
        )
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = wb.active
        rows: list[list[Any]] = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= 200:
                break
            rows.append([str(c) if c is not None else "" for c in row])
        return _truncate(json.dumps(rows, ensure_ascii=False), cap)
    except Exception as e:
        return f"[XLSX parse failed for {filename}: {e}]"
